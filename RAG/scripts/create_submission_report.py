from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[2] / "Bao_cao_he_thong_RAG_HVNH_Hub.docx"
NAVY = "063763"
BLUE = "DCEBFA"
GRAY = "F5F7FA"
BORDER = "D9E2EC"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_border(cell, color: str = BORDER) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for name in ("top", "left", "bottom", "right"):
        edge = borders.find(qn(f"w:{name}"))
        if edge is None:
            edge = OxmlElement(f"w:{name}")
            borders.append(edge)
        edge.set(qn("w:val"), "single")
        edge.set(qn("w:sz"), "6")
        edge.set(qn("w:color"), color)


def set_cell_margins(cell, top=120, start=140, bottom=120, end=140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_text(paragraph, text: str, bold=False, color=None, size=None) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if size:
        run.font.size = Pt(size)
    run.font.name = "Aptos"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    add_text(p, text)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.space_before = Pt(15 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    add_text(p, text, bold=True, color="000000", size=15 if level == 1 else 12)


def add_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.22
    add_text(p, text)


def add_step_table(doc: Document) -> None:
    rows = [
        ("1", "Nguồn chính thức", "Website HVNH, thông báo, PDF, DOCX và tài liệu được khai báo trong cơ sở dữ liệu."),
        ("2", "Crawler", "Tải tài liệu an toàn, tôn trọng robots.txt, giới hạn redirect, timeout và dung lượng."),
        ("3", "Trích xuất", "Đọc HTML, PDF, DOCX; PDF scan được chuyển qua OCR để lấy văn bản."),
        ("4", "Chuẩn hóa và phiên bản", "Lưu nội dung, metadata, hash; tạo phiên bản mới chỉ khi tài liệu thay đổi."),
        ("5", "Chunk và embedding", "Chia tài liệu dài thành đoạn nhỏ, biến từng đoạn thành vector ngữ nghĩa."),
        ("6", "Lập chỉ mục", "Lưu vector vào Qdrant, đồng thời giữ dữ liệu gốc và chunk trong PostgreSQL."),
        ("7", "Hỏi đáp", "Hiểu ngữ cảnh hội thoại, truy xuất hybrid, rerank, sinh trả lời và kiểm tra trích dẫn."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    headers = ["Bước", "Thành phần", "Kết quả"]
    for cell, value in zip(table.rows[0].cells, headers):
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, NAVY)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(p, value, bold=True, color="FFFFFF", size=10)
    for index, (step, name, result) in enumerate(rows):
        cells = table.add_row().cells
        for cell in cells:
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if index % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, GRAY)
        cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_text(cells[0].paragraphs[0], step, bold=True, color=NAVY)
        add_text(cells[1].paragraphs[0], name, bold=True)
        add_text(cells[2].paragraphs[0], result)


def add_technology_table(doc: Document) -> None:
    rows = [
        ("httpx", "Tải website và tài liệu từ Internet."),
        ("BeautifulSoup", "Đọc HTML, lấy nội dung và phát hiện liên kết tài liệu."),
        ("pypdf và python-docx", "Trích xuất text từ PDF có text và DOCX."),
        ("Modal OCR", "Chuyển PDF scan hoặc ảnh thành text có thể tìm kiếm."),
        ("Redis Streams", "Hàng đợi crawl và pipeline; khóa chống xử lý trùng."),
        ("PostgreSQL Neon", "Nguồn dữ liệu chính: tài liệu, phiên bản, metadata, chunk, hội thoại và citation."),
        ("BGE-M3 hoặc Gemini Embedding", "Biến đoạn văn và câu hỏi thành vector thể hiện ngữ nghĩa."),
        ("Qdrant", "Lưu vector và tìm nhanh các đoạn có ý nghĩa gần câu hỏi."),
        ("PostgreSQL Full Text Search", "Tìm chính xác tên người, mã quyết định, số tiền hoặc mã học phần."),
        ("RRF và Heuristic Reranker", "Gộp semantic search với keyword search và xếp lại kết quả."),
        ("Gemini", "Sinh câu trả lời ngắn, sát nguồn và bắt buộc có citation."),
        ("LangGraph RAG Service", "Điều phối hội thoại, truy xuất, sinh trả lời, kiểm chứng và lưu lịch sử."),
    ]
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for cell, value in zip(table.rows[0].cells, ["Công nghệ", "Vai trò trong hệ thống"]):
        set_cell_shading(cell, NAVY)
        set_cell_border(cell, NAVY)
        set_cell_margins(cell)
        add_text(cell.paragraphs[0], value, bold=True, color="FFFFFF", size=10)
    for index, (technology, role) in enumerate(rows):
        cells = table.add_row().cells
        for cell in cells:
            set_cell_border(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if index % 2 == 1:
            for cell in cells:
                set_cell_shading(cell, GRAY)
        add_text(cells[0].paragraphs[0], technology, bold=True, color=NAVY)
        add_text(cells[1].paragraphs[0], role)


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    styles["Normal"].font.size = Pt(11)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    add_text(title, "Báo cáo hệ thống RAG HVNH Hub", bold=True, color="000000", size=24)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(22)
    add_text(subtitle, "Luồng thu thập dữ liệu, lưu trữ tri thức và sinh câu trả lời có căn cứ", color="4F6277", size=12)

    add_heading(doc, "1. Mục tiêu hệ thống")
    add_paragraph(doc, "Hệ thống RAG của HVNH Hub giúp sinh viên hỏi các thông tin chính thức của Học viện Ngân hàng như quy chế, học phí, lịch học, thủ tục, thông báo và sự kiện. Điểm cốt lõi là trợ lý không trả lời theo trí nhớ chung mà phải tìm được tài liệu phù hợp, dùng tài liệu đó làm căn cứ và hiển thị trích dẫn.")
    add_paragraph(doc, "RAG là viết tắt của Retrieval Augmented Generation. Có thể hiểu đơn giản: trước khi AI trả lời, hệ thống đi tìm bằng chứng trong kho tri thức; sau đó AI chỉ diễn đạt lại nội dung đã tìm thấy.")

    add_heading(doc, "2. Luồng tổng quát")
    add_step_table(doc)

    add_heading(doc, "3. Thu thập và xử lý dữ liệu")
    add_heading(doc, "3.1 Quản lý nguồn và lập lịch crawl", 2)
    add_paragraph(doc, "Các URL nguồn được quản lý trong cơ sở dữ liệu thay vì ghi cứng trong code. Scheduler kiểm tra nguồn nào đến lịch crawl, tạo một crawl run và đưa run vào Redis Streams. Worker nhận job từ Redis để xử lý nền. Vì vậy crawl không làm chậm các API đăng nhập, bài viết hoặc chat.")
    add_heading(doc, "3.2 Crawler an toàn", 2)
    add_paragraph(doc, "Crawler dùng httpx để tải dữ liệu. Trước khi tải, hệ thống kiểm tra phạm vi domain, URL công khai, robots.txt, số lần redirect, timeout và kích thước phản hồi. Nếu một URL lỗi, worker ghi nhận lỗi của URL đó nhưng vẫn tiếp tục các URL khác trong phiên crawl.")
    add_heading(doc, "3.3 Trích xuất nội dung", 2)
    add_bullet(doc, "HTML được xử lý bằng BeautifulSoup để lấy text, tiêu đề và liên kết tài liệu.")
    add_bullet(doc, "PDF có text được xử lý bằng pypdf, đồng thời lưu thông tin số trang.")
    add_bullet(doc, "DOCX được đọc bằng python-docx.")
    add_bullet(doc, "PDF scan có quá ít text được chuyển sang Modal OCR để nhận dạng ký tự.")
    add_bullet(doc, "Dữ liệu quá ngắn, không đúng định dạng, hoặc vi phạm chính sách chất lượng được đánh dấu thay vì đưa ngay vào kho hỏi đáp.")

    add_heading(doc, "4. Lưu trữ, phiên bản và lập chỉ mục")
    add_paragraph(doc, "PostgreSQL Neon là nơi lưu dữ liệu gốc và có tính truy vết. Hệ thống lưu nguồn, tài liệu, phiên bản, raw text, metadata, chunk, lịch sử hội thoại, câu trả lời và citation. Mỗi phiên bản có hash nội dung. Khi nguồn không thay đổi, hệ thống tránh tạo dữ liệu trùng lặp; khi có bản mới, hệ thống tạo phiên bản mới.")
    add_paragraph(doc, "Mỗi tài liệu được chia thành chunk khoảng 1.800 ký tự, có overlap khoảng 250 ký tự. Chunk giữ số trang, mục và heading để kết quả có thể quay về đúng phần tài liệu. Sau đó embedding biến mỗi chunk thành vector 1024 chiều. Các vector cùng metadata được lưu vào Qdrant.")
    add_paragraph(doc, "Qdrant không thay PostgreSQL. PostgreSQL giữ sự thật và lịch sử; Qdrant là chỉ mục tìm kiếm nhanh theo ngữ nghĩa. Khi có phiên bản mới, vector của bản cũ được đánh dấu không còn là current. Câu hỏi hiện tại ưu tiên bản mới; câu hỏi có năm học hoặc yêu cầu dữ liệu cũ vẫn có thể truy xuất phiên bản lịch sử.")

    add_heading(doc, "5. Luồng khi người dùng đặt câu hỏi")
    add_heading(doc, "5.1 Giữ ngữ cảnh hội thoại", 2)
    add_paragraph(doc, "ConversationMemoryService lấy các tin nhắn gần nhất và phần tóm tắt lịch sử để viết lại câu hỏi thành một câu độc lập. Ví dụ, sau câu hỏi về giảng viên Nguyễn Thanh Thụy, câu hỏi tiếp theo Thầy dạy những môn nào sẽ được hiểu là Giảng viên Nguyễn Thanh Thụy dạy những môn nào. Câu viết lại mới là câu được đưa vào retrieval.")
    add_heading(doc, "5.2 Phân tích ý định", 2)
    add_paragraph(doc, "Hệ thống phát hiện năm học, loại tài liệu, nhu cầu bản mới nhất và nhu cầu dữ liệu lịch sử. Câu hỏi có từ hiện tại, mới nhất hoặc gần đây sẽ ưu tiên tài liệu current. Câu hỏi có năm học, bản cũ hoặc trước đây sẽ nới điều kiện để tìm đúng phiên bản cũ.")
    add_heading(doc, "5.3 Hybrid retrieval", 2)
    add_paragraph(doc, "Retriever tìm song song theo hai hướng. Semantic search dùng vector trong Qdrant để hiểu ý nghĩa. Lexical search dùng PostgreSQL Full Text Search để giữ độ chính xác với tên riêng, mã quyết định, mã học phần, số tiền và thuật ngữ. Hai danh sách được gộp bằng Reciprocal Rank Fusion, rồi Heuristic Reranker xếp lại theo mức khớp tiêu đề, thực thể, năm học, số liệu, thủ tục và độ mới của tài liệu.")
    add_heading(doc, "5.4 Sinh và kiểm chứng câu trả lời", 2)
    add_paragraph(doc, "Các chunk tốt nhất được đưa cho Gemini dưới dạng context có ID. Gemini phải trả về câu trả lời có citation dạng [1], [2] và danh sách ID context đã sử dụng. Citation Verifier kiểm tra citation có tồn tại, có đúng context hay không và claim có được phần trích dẫn hỗ trợ hay không. Nếu không đủ bằng chứng hoặc citation sai, hệ thống từ chối trả lời thay vì suy diễn.")

    add_heading(doc, "6. Khi nào hệ thống từ chối")
    add_bullet(doc, "Không tìm thấy nguồn chính thức phù hợp trong kho đã lập chỉ mục.")
    add_bullet(doc, "Tên người hoặc thực thể được hỏi không xuất hiện trong các context truy xuất.")
    add_bullet(doc, "Câu hỏi về hôm nay hoặc ngày mai không có tài liệu mới đủ bằng chứng.")
    add_bullet(doc, "LLM không tạo được citation hợp lệ sau số lần thử cho phép.")

    add_heading(doc, "7. Vai trò của các công nghệ")
    add_technology_table(doc)

    add_heading(doc, "8. Kết luận")
    add_paragraph(doc, "Kiến trúc RAG của HVNH Hub được xây dựng theo nguyên tắc có căn cứ và có thể kiểm soát. Crawler đưa dữ liệu chính thức vào kho; PostgreSQL giữ dữ liệu gốc và phiên bản; Qdrant tăng tốc tìm kiếm ngữ nghĩa; Gemini chỉ sinh câu trả lời từ context đã truy xuất; Citation Verifier ngăn câu trả lời không có căn cứ. Nhờ cơ chế current version, historical query và memory hội thoại, hệ thống vừa ưu tiên thông tin mới vừa bảo toàn khả năng trả lời về dữ liệu cũ và các câu hỏi nối tiếp.")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_text(footer, "HVNH Hub - Báo cáo hệ thống RAG", color="62758B", size=9)

    doc.core_properties.title = "Báo cáo hệ thống RAG HVNH Hub"
    doc.core_properties.author = "HVNH Hub"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
