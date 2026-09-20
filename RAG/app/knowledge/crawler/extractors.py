from dataclasses import dataclass, field
from datetime import datetime
from io import BytesIO
import csv
import re
from typing import Any
from urllib.parse import parse_qsl, unquote, urljoin, urlsplit

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


class UnsupportedDocument(ValueError):
    pass


class OCRRequired(UnsupportedDocument):
    """The file is readable as a PDF but contains too little extractable text."""

    pass


@dataclass(frozen=True, slots=True)
class ExtractedSection:
    title: str | None
    content: str
    page: int | None = None


@dataclass(frozen=True, slots=True)
class ExtractedTable:
    name: str | None
    headers: tuple[str, ...]
    rows: tuple[tuple[str, ...], ...]
    page: int | None = None


@dataclass(frozen=True, slots=True)
class ExtractedDocument:
    title: str
    content: str
    links: tuple[str, ...] = ()
    published_at: datetime | None = None
    page_count: int | None = None
    metadata: dict[str, Any] = field(default_factory=dict)
    sections: tuple[ExtractedSection, ...] = ()
    tables: tuple[ExtractedTable, ...] = ()


def _clean(text: str) -> str:
    lines = [re.sub(r"\s+", " ", line).strip() for line in text.replace("\xa0", " ").splitlines()]
    return "\n".join(line for line in lines if line)


_PDF_URL_PATTERN = re.compile(
    r"(?:(?:https?:)?//|/)[^\s\"'<>]+?\.pdf(?:\?[^\s\"'<>]*)?",
    re.IGNORECASE,
)


def _decode_repeatedly(value: str, attempts: int = 3) -> str:
    for _ in range(attempts):
        decoded = unquote(value).replace("\\/", "/")
        if decoded == value:
            break
        value = decoded
    return value


def pdf_text_requires_ocr(content: str) -> bool:
    tokens = re.findall(r"\w+", content, flags=re.UNICODE)
    single_character_ratio = (sum(len(token) == 1 for token in tokens) / len(tokens)) if tokens else 1.0
    return len(content) < 40 or (len(tokens) >= 50 and single_character_ratio > 0.55)


def _discover_html_links(soup: BeautifulSoup, body: bytes, base_url: str) -> tuple[str, ...]:
    """Discover navigational and embedded document URLs, including HVNH PDF viewers."""
    candidates: list[str] = []
    for node in soup.select("a[href]"):
        if node.get("href"):
            candidates.append(node["href"])
    for selector, attributes in (
        ("iframe", ("src", "data-src", "data-url")),
        ("object", ("data", "data-src", "data-url")),
        ("embed", ("src", "data-src", "data-url")),
        ("source", ("src", "data-src")),
    ):
        for node in soup.select(selector):
            candidates.extend(node.get(attribute) for attribute in attributes if node.get(attribute))

    # Some HVNH pages create the viewer from JavaScript instead of a DOM element.
    raw_html = body.decode("utf-8", errors="replace").replace("\\/", "/")
    candidates.extend(match.group(0) for match in _PDF_URL_PATTERN.finditer(raw_html))

    discovered: list[str] = []
    for raw in candidates:
        absolute = urljoin(base_url, str(raw).strip())
        if not absolute:
            continue
        discovered.append(absolute)
        for key, value in parse_qsl(urlsplit(absolute).query, keep_blank_values=False):
            if key.casefold() not in {"pdfurl", "pdf", "fileurl", "documenturl"}:
                continue
            document_url = _decode_repeatedly(value).strip()
            if ".pdf" in document_url.casefold():
                discovered.append(urljoin(absolute, document_url))
    return tuple(dict.fromkeys(discovered))


def _html(body: bytes, url: str) -> ExtractedDocument:
    soup = BeautifulSoup(body, "html.parser")
    for node in soup.select("script,style,noscript,svg,form,nav,header,footer,aside"):
        node.decompose()
    title_node = soup.select_one("meta[property='og:title']")
    title = (title_node.get("content") if title_node else None) or (soup.h1.get_text(" ", strip=True) if soup.h1 else None) or (soup.title.get_text(" ", strip=True) if soup.title else url)
    path = urlsplit(url).path.casefold()
    if "/gioi_thieu/giang_vien/" in path:
        main = soup.select_one(".stu-db .container.pg-inn") or soup.select_one(".stu-db")
    else:
        main = soup.select_one(".tin-chi-tiet, article, main, [itemprop='articleBody'], .news-detail, .detail-content, .news-content, .content-detail, .stu-db, .ho-event.pg-eve-main.pg-blog")
    main = main or soup.body or soup
    links = _discover_html_links(soup, body, url)
    published_at = None
    date_node = soup.select_one("meta[name='pubdate'], meta[itemprop='datePublished'], meta[property='article:published_time'], meta[name='date'], time[datetime]")
    raw_date = (date_node.get("content") or date_node.get("datetime")) if date_node else None
    if raw_date:
        try: published_at = datetime.fromisoformat(str(raw_date).strip().replace("Z", "+00:00"))
        except ValueError: pass
    content = _clean(main.get_text("\n", strip=True))
    sections = tuple(ExtractedSection(_clean(node.get_text(" ", strip=True)), content)
                     for node in main.select("h1,h2,h3")[:50])
    return ExtractedDocument(_clean(str(title)), content, links, published_at,
                             sections=sections or (ExtractedSection(_clean(str(title)), content),))


def _pdf(body: bytes, url: str) -> ExtractedDocument:
    reader = PdfReader(BytesIO(body))
    pages = [_clean(page.extract_text() or "") for page in reader.pages]
    offset, ranges = 0, []
    for page_number, page in enumerate(pages, 1):
        if page_number > 1: offset += 2
        start = offset; offset += len(page); ranges.append((page_number, start, offset))
    title = str(reader.metadata.title) if reader.metadata and reader.metadata.title else unquote(url.rsplit("/", 1)[-1])
    content = "\n\n".join(pages)
    if pdf_text_requires_ocr(content):
        raise OCRRequired("PDF has too little extractable text and requires OCR")
    sections = tuple(ExtractedSection(f"Trang {number}", page, number)
                     for number, page in enumerate(pages, 1) if page)
    return ExtractedDocument(_clean(title), content, page_count=len(pages),
                             metadata={"page_ranges": ranges}, sections=sections)


def _docx(body: bytes, url: str) -> ExtractedDocument:
    doc = Document(BytesIO(body))
    blocks = [paragraph.text for paragraph in doc.paragraphs]
    blocks.extend(" | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows)
    title = next((p.text for p in doc.paragraphs if p.style and p.style.name.startswith("Title") and p.text.strip()), url.rsplit("/", 1)[-1])
    sections = tuple(ExtractedSection(_clean(p.text), _clean(p.text)) for p in doc.paragraphs
                     if p.text.strip() and p.style and p.style.name.startswith("Heading"))
    tables = []
    for number, table in enumerate(doc.tables, 1):
        rows = tuple(tuple(_clean(cell.text) for cell in row.cells) for row in table.rows)
        if rows:
            tables.append(ExtractedTable(f"Table {number}", rows[0], rows[1:]))
    return ExtractedDocument(_clean(title), _clean("\n".join(blocks)), sections=sections,
                             tables=tuple(tables))


def _xlsx(body: bytes, url: str) -> ExtractedDocument:
    workbook = load_workbook(BytesIO(body), read_only=True, data_only=True)
    lines: list[str] = []
    tables: list[ExtractedTable] = []
    for sheet in workbook.worksheets[:20]:
        lines.append(f"Sheet: {sheet.title}")
        sheet_rows: list[tuple[str, ...]] = []
        for row_number, row in enumerate(sheet.iter_rows(values_only=True)):
            if row_number >= 10000:
                break
            values = [str(value) for value in row if value is not None]
            if values:
                lines.append(" | ".join(values))
                sheet_rows.append(tuple(values))
        if sheet_rows:
            tables.append(ExtractedTable(sheet.title, sheet_rows[0], tuple(sheet_rows[1:])))
    workbook.close()
    return ExtractedDocument(url.rsplit("/", 1)[-1], _clean("\n".join(lines)), tables=tuple(tables))


def _csv(body: bytes, url: str) -> ExtractedDocument:
    decoded = body.decode("utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(decoded[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    lines: list[str] = []
    parsed_rows: list[tuple[str, ...]] = []
    for row_number, row in enumerate(csv.reader(decoded.splitlines(), dialect), 1):
        if row_number > 10000:
            break
        values = [value.strip() for value in row if value.strip()]
        if values:
            lines.append(" | ".join(values))
            parsed_rows.append(tuple(values))
    tables = (ExtractedTable(unquote(url.rsplit("/", 1)[-1]), parsed_rows[0], tuple(parsed_rows[1:])),) if parsed_rows else ()
    return ExtractedDocument(unquote(url.rsplit("/", 1)[-1]), _clean("\n".join(lines)), tables=tables)


def _sniff_content_type(body: bytes, declared: str) -> str:
    """Recover from attachment responses incorrectly sent as octet-stream."""
    prefix = body[:512].lstrip().lower()
    if body.startswith(b"%PDF-"):
        return "application/pdf"
    if prefix.startswith((b"<!doctype html", b"<html", b"<?xml")) and b"html" in prefix:
        return "text/html"
    return declared


def extract_document(body: bytes, content_type: str, url: str, *, allow_short: bool = False) -> ExtractedDocument:
    lower_path = url.lower().split("?", 1)[0]
    content_type = _sniff_content_type(body, content_type)
    if content_type in {"text/html", "application/xhtml+xml"} or lower_path.endswith(("/", ".html", ".htm")):
        result = _html(body, url)
    elif content_type == "application/pdf" or lower_path.endswith(".pdf"):
        result = _pdf(body, url)
    elif content_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document" or lower_path.endswith(".docx"):
        result = _docx(body, url)
    elif content_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" or lower_path.endswith(".xlsx"):
        result = _xlsx(body, url)
    elif content_type in {"text/csv", "application/csv"} or lower_path.endswith((".csv", ".tsv")):
        result = _csv(body, url)
    elif content_type.startswith("text/plain"):
        result = ExtractedDocument(url.rsplit("/", 1)[-1], _clean(body.decode("utf-8", errors="replace")))
    else:
        raise UnsupportedDocument(f"unsupported content type: {content_type or 'unknown'}")
    if len(result.content) < 40 and not allow_short:
        raise UnsupportedDocument("document has too little extractable text")
    return result
