from io import BytesIO

import pytest
from pypdf import PdfWriter
from docx import Document
from openpyxl import Workbook

from app.knowledge.crawler.extractors import OCRRequired, extract_document, pdf_text_requires_ocr
from app.knowledge.crawler.security import UnsafeURL, canonicalize_url, is_in_scope
from app.knowledge.crawler.metadata import infer_academic_year, infer_document_type, infer_metadata, infer_published_at
from app.knowledge.crawler.worker import _resume_checkpoint, _should_store_page
from app.models.crawler import AICrawlRun
from app.models.crawler import AISource
from app.models.enums import AIDocumentType
from app.knowledge.quality import assess_extraction


def test_canonical_url_removes_tracking_and_fragment() -> None:
    assert canonicalize_url("https://HVNH.EDU.VN/a?utm_source=x&id=2#top") == "https://hvnh.edu.vn/a?id=2"


def test_url_scope_does_not_accept_domain_suffix_attack() -> None:
    assert is_in_scope("https://hvnh.edu.vn/sinh-vien/a", ["hvnh.edu.vn"], ["/sinh-vien"])
    assert not is_in_scope("https://hvnh.edu.vn.attacker.test/sinh-vien", ["hvnh.edu.vn"], [])


def test_canonical_url_rejects_http_and_credentials() -> None:
    with pytest.raises(UnsafeURL): canonicalize_url("http://hvnh.edu.vn")
    with pytest.raises(UnsafeURL): canonicalize_url("https://user:pass@hvnh.edu.vn")


def test_html_extractor_removes_navigation_and_finds_links() -> None:
    body = b"<html><head><title>Notice</title></head><body><nav>Menu</nav><article><h1>Notice</h1><p>" + b"Official content " * 5 + b"</p><a href='/file.pdf'>PDF</a></article></body></html>"
    result = extract_document(body, "text/html", "https://hvnh.edu.vn/news/1")
    assert "Menu" not in result.content
    assert "Official content" in result.content
    assert "https://hvnh.edu.vn/file.pdf" in result.links


def test_html_extractor_discovers_pdf_inside_hvnh_viewer() -> None:
    body = (b"<html><body><article><p>Official tuition information with enough text.</p>"
            b"<iframe src='/medias/pdfviewer/viewer/viewer.aspx?pdfurl="
            b"https%3A%2F%2Fhvnh.edu.vn%2Fmedias%2Ftckt%2Ffee.pdf'></iframe>"
            b"</article></body></html>")
    result = extract_document(body, "text/html", "https://hvnh.edu.vn/tckt/vi/hoc-phi/item.html")
    assert "https://hvnh.edu.vn/medias/tckt/fee.pdf" in result.links


def test_html_extractor_discovers_viewer_pdf_from_anchor_and_lazy_embed() -> None:
    body = (
        b"<html><body><article><p>Official document information with enough text.</p>"
        b"<a href='/viewer.aspx?PDFURL=%252Fmedias%252Frules%252Fregulation.pdf'>Open</a>"
        b"<iframe data-src='/medias/fees/tuition.pdf'></iframe>"
        b"</article></body></html>"
    )
    result = extract_document(body, "text/html", "https://hvnh.edu.vn/news/item.html")
    assert "https://hvnh.edu.vn/medias/rules/regulation.pdf" in result.links
    assert "https://hvnh.edu.vn/medias/fees/tuition.pdf" in result.links


def test_html_extractor_discovers_pdf_created_by_javascript() -> None:
    body = (
        b"<html><body><article><p>Official document information with enough text.</p></article>"
        b"<script>const documentUrl = 'https://hvnh.edu.vn/medias/rules/rule.pdf';</script>"
        b"</body></html>"
    )
    result = extract_document(body, "text/html", "https://hvnh.edu.vn/news/item.html")
    assert "https://hvnh.edu.vn/medias/rules/rule.pdf" in result.links


def test_image_only_pdf_is_marked_for_ocr() -> None:
    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(output)
    with pytest.raises(OCRRequired):
        extract_document(output.getvalue(), "application/pdf", "https://hvnh.edu.vn/scan.pdf")


def test_pdf_is_detected_when_server_returns_octet_stream() -> None:
    output = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=100, height=100)
    writer.write(output)
    with pytest.raises(OCRRequired):
        extract_document(output.getvalue(), "application/octet-stream",
                         "https://hvnh.edu.vn/download?id=123")


def test_csv_extractor_supports_semicolon_delimiter_and_utf8_bom() -> None:
    body = "\ufeffHọc phần;Số tín chỉ\nTrí tuệ nhân tạo;3\n".encode("utf-8")
    result = extract_document(body, "text/csv", "https://hvnh.edu.vn/data.csv")
    assert "Trí tuệ nhân tạo | 3" in result.content
    assert result.tables[0].headers == ("Học phần", "Số tín chỉ")


def test_docx_extractor_preserves_table_structure() -> None:
    output = BytesIO(); document = Document(); document.add_heading("Quy chế", level=1)
    document.add_paragraph("Nội dung chính thức của Học viện Ngân hàng đủ dài để kiểm thử.")
    table = document.add_table(rows=2, cols=2); table.cell(0, 0).text = "Hệ"; table.cell(0, 1).text = "Mức thu"
    table.cell(1, 0).text = "Chuẩn"; table.cell(1, 1).text = "100"
    document.save(output)
    result = extract_document(output.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                              "https://hvnh.edu.vn/quy-che.docx")
    assert result.tables[0].headers == ("Hệ", "Mức thu")
    assert result.tables[0].rows == (("Chuẩn", "100"),)


def test_xlsx_extractor_preserves_sheet_and_headers() -> None:
    output = BytesIO(); workbook = Workbook(); sheet = workbook.active; sheet.title = "Học phí"
    sheet.append(["Chương trình", "Mức thu"]); sheet.append(["Chuẩn", 100]); workbook.save(output)
    result = extract_document(output.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                              "https://hvnh.edu.vn/hoc-phi.xlsx")
    assert result.tables[0].name == "Học phí"
    assert result.tables[0].rows == (("Chuẩn", "100"),)


def test_quality_gate_rejects_corrupt_extraction() -> None:
    result = extract_document(("� !!! " * 30).encode(), "text/plain", "https://hvnh.edu.vn/bad.txt")
    quality = assess_extraction(result)
    assert quality.accepted is False
    assert quality.code == "LOW_EXTRACTION_QUALITY"


def test_low_quality_spaced_pdf_text_heuristic() -> None:
    assert pdf_text_requires_ocr(" ".join("A" for _ in range(60)))
    assert not pdf_text_requires_ocr("Mức học phí là 680.000 đồng trên một tín chỉ. " * 10)


def test_itde_profile_extractor_prefers_profile_container() -> None:
    body = (b"<html><head><title>Lecturer</title></head><body><div class='menu'>Noisy menu "
            b"content repeated many times</div><div class='stu-db'><h1>Lecturer Name</h1>"
            b"<p>Official lecturer profile and teaching subjects.</p></div></body></html>")
    result = extract_document(body, "text/html", "https://itde.hvnh.edu.vn/gioi_thieu/giang_vien/1.html")
    assert "Official lecturer profile" in result.content
    assert "Noisy menu" not in result.content


def test_itde_profile_extractor_uses_inner_profile_container() -> None:
    body = (b"<html><head><title>Lecturer</title></head><body><section class='stu-db'>"
            b"<div>Navigation repeated outside the profile</div><div class='container pg-inn'>"
            b"<h4>Lecturer Name</h4><p>Department and teaching subjects with enough official text.</p>"
            b"</div></section></body></html>")
    result = extract_document(body, "text/html",
                              "https://itde.hvnh.edu.vn/gioi_thieu/giang_vien/1.html")
    assert "teaching subjects" in result.content
    assert "Navigation repeated" not in result.content


def test_store_configured_extensionless_html_detail_page() -> None:
    source = AISource(
        name="Online HVNH",
        base_url="https://online.hvnh.edu.vn/News/Type/1017",
        allowed_domains=["online.hvnh.edu.vn"],
        allowed_path_prefixes=["/News/Type/1017", "/News/Detail/"],
        metadata_={"html_detail_path_prefixes": ["/News/Detail/"]},
    )
    assert _should_store_page(source, "https://online.hvnh.edu.vn/News/Detail/123", "text/html")
    assert not _should_store_page(source, "https://online.hvnh.edu.vn/News/Type/1017", "text/html")
    assert _should_store_page(source, "https://online.hvnh.edu.vn/files/quy-che.pdf", "application/pdf")


def test_short_listing_can_be_used_only_for_link_discovery() -> None:
    body = b"<html><body><a href='/News/Detail/1'>Open</a></body></html>"
    with pytest.raises(ValueError):
        extract_document(body, "text/html", "https://online.hvnh.edu.vn/News/Type/1017")
    result = extract_document(
        body, "text/html", "https://online.hvnh.edu.vn/News/Type/1017", allow_short=True
    )
    assert result.links == ("https://online.hvnh.edu.vn/News/Detail/1",)


def test_checkpoint_is_resumed_only_when_immediately_previous_run_is_partial() -> None:
    frontier = [{"url": "https://hvnh.edu.vn/next", "depth": 1, "parent": None}]
    partial = AICrawlRun(status="PARTIAL", metadata_={"frontier": frontier, "visited": ["seed"]})
    succeeded = AICrawlRun(status="SUCCEEDED", metadata_={"frontier": frontier})
    assert _resume_checkpoint(partial, fresh_start=False)["frontier"] == frontier
    assert _resume_checkpoint(succeeded, fresh_start=False) == {}
    assert _resume_checkpoint(partial, fresh_start=True) == {}


def test_partial_without_remaining_frontier_starts_clean_retry() -> None:
    partial = AICrawlRun(status="PARTIAL", metadata_={"frontier": [], "visited": ["seed"]})
    assert _resume_checkpoint(partial, fresh_start=False) == {}


def test_metadata_does_not_take_unlabelled_year_from_navigation() -> None:
    metadata = infer_metadata(
        "Giảng viên | Triệu Thu Hương",
        "Thông tin giảng viên\nBài viết mới: tuyển sinh 2026/2027",
    )
    assert metadata.academic_year is None


def test_metadata_accepts_explicit_academic_year_from_body() -> None:
    metadata = infer_metadata("Thông báo học phí", "Áp dụng trong năm học 2025/2026.")
    assert metadata.academic_year == "2025/2026"


def test_hvnh_metadata_inference() -> None:
    assert infer_academic_year("năm học 2026–2027") == "2026/2027"
    assert infer_academic_year("hoc-ky-20262027") == "2026/2027"
    assert infer_document_type("Thông báo tuyển sinh") == AIDocumentType.ANNOUNCEMENT
    assert infer_published_at("Ngày 17/07/2026").isoformat() == "2026-07-17T00:00:00+07:00"
    assert infer_published_at("Kỷ niệm 26/3/1931, đăng Ngày 20/03/2026").year == 2026
    assert infer_published_at("Ngày đăng: 31/07/2023").isoformat() == "2023-07-31T00:00:00+07:00"
    assert infer_published_at("Đăng ngày 01-06-2026").isoformat() == "2026-06-01T00:00:00+07:00"
    assert infer_published_at("Hà Nội, ngày 28 tháng 7 năm 2025").isoformat() == "2025-07-28T00:00:00+07:00"
