import asyncio
import io

from docx import Document

from src.services.study_document_ai import LLMFileResult, _assess_file, aggregate_file_results, extract_text
from src.services.study_moderation_llm import StudyModerationLLM


def test_bundle_decision_uses_per_file_results():
    good = {"file_name": "slide.pdf", "recommendation": "APPROVE"}
    review = {"file_name": "scan.pdf", "recommendation": "REVIEW"}
    bad = {"file_name": "spam.pdf", "recommendation": "REJECT"}
    assert aggregate_file_results([good])[0] == "APPROVE"
    assert aggregate_file_results([good, review])[0] == "REVIEW"
    assert aggregate_file_results([good, bad])[0] == "REJECT"
    assert aggregate_file_results([])[0] == "REVIEW"


def test_docx_extraction_includes_tables_and_unsupported_is_review():
    document = Document()
    document.add_paragraph("Mô hình dữ liệu quan hệ")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Khóa chính"
    stream = io.BytesIO()
    document.save(stream)
    sample, limited = extract_text(stream.getvalue(), "docx")
    assert "Mô hình dữ liệu quan hệ" in sample
    assert "Khóa chính" in sample
    assert limited is False
    assert extract_text(b"not an OCR image", "jpg") == ("", True)


def test_unverifiable_model_quote_falls_back_to_review():
    class FakeProvider:
        async def generate_structured(self, *args, **kwargs):
            return {
                "subject_relevance": "RELEVANT", "content_type": "ACADEMIC",
                "evidence_quote": "Câu trích không hề nằm trong file",
                "reason": "Phù hợp với môn học.",
            }

    file = {"id": "file-1", "original_name": "notes.pdf"}
    subject = {"major_name": "CNTT", "subject_code": "IS21A", "subject_name": "Cơ sở dữ liệu"}
    result = asyncio.run(_assess_file(FakeProvider(), file, subject,
        "Mô hình quan hệ và bài tập truy vấn SQL. " * 10, False))
    assert result["recommendation"] == "REVIEW"
    assert result["error_code"] == "UNVERIFIED_EVIDENCE"


def test_verified_file_evidence_is_advisory_and_limited_sample_needs_review():
    quote = "Bài tập truy vấn SQL trên mô hình quan hệ"

    class FakeProvider:
        async def generate_structured(self, *args, **kwargs):
            return {
                "subject_relevance": "RELEVANT", "content_type": "ACADEMIC",
                "evidence_quote": quote, "reason": "Nội dung phù hợp với môn học.",
            }

    file = {"id": "file-2", "original_name": "sql.pdf"}
    subject = {"major_name": "CNTT", "subject_code": "IS21A", "subject_name": "Cơ sở dữ liệu"}
    sample = f"[Trang/phần 2]\n{quote}. " + "Mô hình dữ liệu quan hệ. " * 10
    approved = asyncio.run(_assess_file(FakeProvider(), file, subject, sample, False))
    assert approved["recommendation"] == "APPROVE"
    assert approved["evidence_location"] == "Trang/phần 2"
    limited = asyncio.run(_assess_file(FakeProvider(), file, subject, sample, True))
    assert limited["recommendation"] == "REVIEW"


def test_moderation_uses_its_own_key_and_model(monkeypatch):
    from src.services import study_moderation_llm

    observed = {}

    class FakeModels:
        async def generate_content(self, *, model, contents, config):
            observed["model"] = model
            observed["contents"] = contents
            return type("Response", (), {"parsed": {
                "subject_relevance": "RELEVANT", "content_type": "ACADEMIC",
                "evidence_quote": "A verified quote", "reason": "Study material",
            }})()

    class FakeClient:
        def __init__(self, *, api_key):
            observed["api_key"] = api_key
            self.aio = type("AsyncClient", (), {"models": FakeModels()})()

    monkeypatch.setattr(study_moderation_llm.settings, "STUDY_MODERATION_API_KEY", "moderation-only-key")
    monkeypatch.setattr(study_moderation_llm.settings, "STUDY_MODERATION_MODEL", "moderation-model")
    monkeypatch.setattr(study_moderation_llm.genai, "Client", FakeClient)

    provider = StudyModerationLLM()
    result = asyncio.run(provider.generate_structured("Classify", "document sample", LLMFileResult))
    assert result.subject_relevance == "RELEVANT"
    assert observed == {"api_key": "moderation-only-key", "model": "moderation-model",
                        "contents": "document sample"}
