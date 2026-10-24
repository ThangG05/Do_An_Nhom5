import uuid
from io import BytesIO
from datetime import UTC, date, datetime, time, timedelta
from zoneinfo import ZoneInfo

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy import bindparam, func, or_, select, text
from sqlalchemy.orm import Session

from src.db.models.user import AccountStatus, SystemRole, User
from src.db.models.media import MediaFile, MediaStatus
from src.db.models.post import Post, PostStatus, PostVisibility
from src.db.models.user import Profile
from src.models.admin import (AdminDashboardResponse, AdminUserPage, AdminUserResponse,
                              AdminProfilePreview, AdminPublicPostPage, AuditLogResponse, DisciplineRequest, GroupAdminInfoResponse,
                              RAGUserUsageResponse)
from src.services.storage_service import create_download_url
from src.services import post_service

DEDICATED_ADMIN_SLUGS = {"pass-do", "ghep-phong-tim-tro", "su-kien", "hoc-tap"}


def profile_preview(db: Session, user_id: uuid.UUID) -> AdminProfilePreview:
    user = db.get(User, user_id)
    if not user or user.deleted_at is not None:
        raise HTTPException(404, "Không tìm thấy tài khoản.")
    profile = user.profile
    media_ids = [media_id for media_id in (profile.avatar_media_id, profile.cover_media_id) if media_id] if profile else []
    media = {item.id: item for item in db.scalars(select(MediaFile).where(
        MediaFile.id.in_(media_ids), MediaFile.status == MediaStatus.READY,
        MediaFile.deleted_at.is_(None),
    ))} if media_ids else {}
    return AdminProfilePreview(
        id=str(user.id), name=profile.full_name if profile else user.username,
        username=user.username,
        avatar=create_download_url(media[profile.avatar_media_id]) if profile and profile.avatar_media_id in media else None,
        cover=create_download_url(media[profile.cover_media_id]) if profile and profile.cover_media_id in media else None,
        bio=profile.bio or "" if profile else "",
        student_code=profile.student_code if profile else None,
        faculty=profile.faculty if profile else None,
        joined_at=user.created_at.isoformat(), status=user.status.value,
    )


def public_posts(db: Session, admin: User, limit: int, offset: int, query: str = "", status: str = "ALL", author_id: uuid.UUID | None = None) -> AdminPublicPostPage:
    filters = [Post.group_id.is_(None), Post.visibility == PostVisibility.PUBLIC,
               Post.deleted_at.is_(None), Post.status.in_((PostStatus.APPROVED, PostStatus.HIDDEN))]
    if status != "ALL":
        filters.append(Post.status == PostStatus(status))
    if author_id:
        filters.append(Post.author_id == author_id)
    term = query.strip()
    if term:
        matched_authors = select(User.id).outerjoin(Profile, Profile.user_id == User.id).where(
            or_(User.username.ilike(f"%{term}%"), Profile.full_name.ilike(f"%{term}%")))
        filters.append(or_(Post.content.ilike(f"%{term}%"), Post.author_id.in_(matched_authors)))
    total = db.scalar(select(func.count()).select_from(Post).where(*filters)) or 0
    posts = list(db.scalars(select(Post).where(*filters).order_by(Post.created_at.desc(), Post.id.desc()).limit(limit).offset(offset)))
    return AdminPublicPostPage(items=post_service._serialize_feed_posts(db, posts, admin.id), total=total, limit=limit, offset=offset)


def moderate_public_post(db: Session, post_id: uuid.UUID, admin: User, action: str, reason: str) -> None:
    post = db.scalar(select(Post).where(Post.id == post_id).with_for_update())
    if not post or post.group_id is not None or post.visibility != PostVisibility.PUBLIC or post.deleted_at is not None:
        raise HTTPException(404, "Không tìm thấy bài viết Feed công khai.")
    expected = PostStatus.APPROVED if action == "HIDE" else PostStatus.HIDDEN
    if post.status != expected:
        raise HTTPException(409, "Trạng thái bài viết đã thay đổi. Hãy tải lại danh sách.")
    post.status = PostStatus.HIDDEN if action == "HIDE" else PostStatus.APPROVED
    post.reviewed_by = admin.id
    post.reviewed_at = datetime.now(UTC)
    post.rejection_reason = reason.strip() if action == "HIDE" else None
    db.execute(text("INSERT INTO audit_logs(actor_id,action,target_type,target_id,metadata) VALUES(:actor,:action,'POST',:target,CAST(:metadata AS jsonb))"),
               {"actor": admin.id, "action": f"PUBLIC_POST_{action}", "target": post.id,
                "metadata": __import__('json').dumps({"reason": reason.strip()}, ensure_ascii=False)})
    db.commit()


def group_admins(db: Session, group_id: uuid.UUID) -> list[GroupAdminInfoResponse]:
    exists = db.execute(text("SELECT 1 FROM groups WHERE id=:gid AND status='ACTIVE'"),
                        {"gid": group_id}).first()
    if not exists:
        raise HTTPException(404, "Không tìm thấy nhóm.")
    rows = db.execute(text("""
        SELECT u.id, u.email, u.username, u.status::text AS status,
               COALESCE(p.full_name, u.username) AS full_name,
               p.student_code, p.faculty, p.avatar_media_id
        FROM group_members gm
        JOIN users u ON u.id=gm.user_id
        LEFT JOIN profiles p ON p.user_id=u.id
        WHERE gm.group_id=:gid AND gm.role='ADMIN' AND u.deleted_at IS NULL
        ORDER BY lower(COALESCE(p.full_name, u.username)), u.id
    """), {"gid": group_id}).mappings().all()
    avatar_ids = {row["avatar_media_id"] for row in rows if row["avatar_media_id"]}
    avatars = {media.id: media for media in db.scalars(select(MediaFile).where(
        MediaFile.id.in_(avatar_ids), MediaFile.status == MediaStatus.READY,
        MediaFile.deleted_at.is_(None),
    )).all()} if avatar_ids else {}
    return [GroupAdminInfoResponse(
        id=row["id"], full_name=row["full_name"], email=row["email"],
        student_code=row["student_code"], faculty=row["faculty"], status=row["status"],
        avatar_url=create_download_url(avatars[row["avatar_media_id"]])
        if row["avatar_media_id"] in avatars else None,
    ) for row in rows]


def _count(db: Session, sql: str, params: dict | None = None) -> int:
    return int(db.execute(text(sql), params or {}).scalar() or 0)


def dashboard(db: Session) -> AdminDashboardResponse:
    totals = db.execute(text("""
        SELECT
          (SELECT count(*) FROM users WHERE deleted_at IS NULL) AS total_users,
          (SELECT count(*) FROM users WHERE status='ACTIVE' AND deleted_at IS NULL) AS active_users,
          (SELECT count(*) FROM users WHERE status='LOCKED' AND deleted_at IS NULL) AS locked_users,
          (SELECT count(*) FROM users WHERE status='DISABLED' AND deleted_at IS NULL) AS disabled_users,
          (SELECT count(*) FROM posts WHERE status='PENDING' AND deleted_at IS NULL) AS pending_posts,
          (SELECT count(*) FROM posts WHERE status='APPROVED' AND deleted_at IS NULL) AS approved_posts,
          (SELECT count(*) FROM posts WHERE status='REJECTED' AND deleted_at IS NULL) AS rejected_posts,
          (SELECT count(*) FROM post_likes) AS total_likes,
          (SELECT count(*) FROM comments WHERE deleted_at IS NULL) AS total_comments,
          (SELECT count(*) FROM messages WHERE deleted_at IS NULL) AS total_messages
    """)).mappings().one()
    group_rows = db.execute(text("SELECT g.id,g.name,g.slug,count(DISTINCT gm.user_id) members,count(DISTINCT p.id) FILTER(WHERE p.status='APPROVED' AND p.deleted_at IS NULL) approved_posts,count(DISTINCT p.id) FILTER(WHERE p.status='PENDING' AND p.deleted_at IS NULL) pending_posts FROM groups g LEFT JOIN group_members gm ON gm.group_id=g.id LEFT JOIN posts p ON p.group_id=g.id WHERE g.status='ACTIVE' GROUP BY g.id ORDER BY g.name")).mappings()
    activity_rows = db.execute(text("""
        SELECT day::date AS date,
          (SELECT count(*) FROM users u WHERE u.created_at >= day AND u.created_at < day + interval '1 day' AND u.deleted_at IS NULL) AS new_users,
          (SELECT count(*) FROM posts p WHERE p.created_at >= day AND p.created_at < day + interval '1 day' AND p.deleted_at IS NULL) AS posts,
          (SELECT count(*) FROM messages m WHERE m.created_at >= day AND m.created_at < day + interval '1 day' AND m.deleted_at IS NULL) AS messages
        FROM generate_series(current_date - interval '13 days', current_date, interval '1 day') day
        ORDER BY day
    """)).mappings()
    return AdminDashboardResponse(
        total_users=totals["total_users"], active_users=totals["active_users"],
        locked_users=totals["locked_users"], disabled_users=totals["disabled_users"],
        pending_posts=totals["pending_posts"], approved_posts=totals["approved_posts"],
        rejected_posts=totals["rejected_posts"], total_likes=totals["total_likes"],
        total_comments=totals["total_comments"], total_messages=totals["total_messages"],
        groups=[dict(row) for row in group_rows],
        activity=[{"date":row["date"].isoformat(),"new_users":row["new_users"],"posts":row["posts"],"messages":row["messages"]} for row in activity_rows],
    )


def export_workbook(db: Session) -> BytesIO:
    stats = dashboard(db)
    workbook = Workbook()
    overview = workbook.active
    overview.title = "Tổng quan"
    overview.append(["BÁO CÁO VẬN HÀNH HVNH HUB", datetime.now(UTC).strftime("%d/%m/%Y %H:%M UTC")])
    overview.append([])
    overview.append(["Chỉ số", "Giá trị"])
    metrics = [
        ("Tổng người dùng", stats.total_users), ("Người dùng hoạt động", stats.active_users),
        ("Tài khoản bị khóa", stats.locked_users), ("Tài khoản vô hiệu hóa", stats.disabled_users),
        ("Bài chờ duyệt", stats.pending_posts), ("Bài đã duyệt", stats.approved_posts),
        ("Bài bị từ chối", stats.rejected_posts), ("Lượt thích", stats.total_likes),
        ("Bình luận", stats.total_comments), ("Tin nhắn", stats.total_messages),
    ]
    for row in metrics: overview.append(row)

    groups_sheet = workbook.create_sheet("Nhóm")
    groups_sheet.append(["Tên nhóm", "Slug", "Thành viên", "Bài đã duyệt", "Bài chờ duyệt"])
    for group in stats.groups:
        groups_sheet.append([group["name"], group["slug"], group["members"], group["approved_posts"], group["pending_posts"]])

    activity_sheet = workbook.create_sheet("Hoạt động 14 ngày")
    activity_sheet.append(["Ngày", "Người dùng mới", "Bài viết", "Tin nhắn"])
    for item in stats.activity:
        activity_sheet.append([item["date"], item["new_users"], item["posts"], item["messages"]])

    users_sheet = workbook.create_sheet("Người dùng")
    users_sheet.append(["Email", "Username", "Họ tên", "Mã sinh viên", "Khoa", "Trạng thái", "Vai trò", "Ngày tạo", "Đăng nhập gần nhất"])
    rows = db.execute(text("""
        SELECT u.email,u.username,COALESCE(p.full_name,u.username) full_name,p.student_code,p.faculty,
               u.status::text status,u.system_role::text system_role,u.created_at,u.last_login_at
        FROM users u LEFT JOIN profiles p ON p.user_id=u.id
        WHERE u.deleted_at IS NULL ORDER BY u.created_at DESC
    """)).mappings()
    for row in rows:
        users_sheet.append([row["email"],row["username"],row["full_name"],row["student_code"],row["faculty"],row["status"],row["system_role"],row["created_at"].replace(tzinfo=None),row["last_login_at"].replace(tzinfo=None) if row["last_login_at"] else None])

    reports_sheet=workbook.create_sheet("Báo cáo vi phạm")
    reports_sheet.append(["Loại","Đối tượng","Lý do","Trạng thái","Ngày tạo","Ngày xử lý"])
    report_rows=db.execute(text("SELECT target_type::text,target_id,reason,status::text,created_at,handled_at FROM reports ORDER BY created_at DESC")).mappings()
    for row in report_rows:reports_sheet.append([row["target_type"],str(row["target_id"]),row["reason"],row["status"],row["created_at"].replace(tzinfo=None),row["handled_at"].replace(tzinfo=None) if row["handled_at"] else None])

    navy = PatternFill("solid", fgColor="002855")
    for sheet in workbook.worksheets:
        header_row = 3 if sheet is overview else 1
        for cell in sheet[header_row]:
            cell.fill = navy
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(horizontal="center")
        sheet.freeze_panes = f"A{header_row + 1}"
        sheet.auto_filter.ref = sheet.dimensions if sheet.max_row >= header_row else None
        for column in sheet.columns:
            letter = column[0].column_letter
            sheet.column_dimensions[letter].width = min(42, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
    overview["A1"].font = Font(size=16, bold=True, color="002855")
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


def _serialize_user(row, groups: list[dict], usage: dict | None = None) -> AdminUserResponse:
    usage = usage or {}
    return AdminUserResponse(id=str(row["id"]),email=row["email"],username=row["username"],full_name=row["full_name"],student_code=row["student_code"],faculty=row["faculty"],status=str(row["status"]),system_role=str(row["system_role"]),admin_groups=groups,created_at=row["created_at"].isoformat(),last_login_at=row["last_login_at"].isoformat() if row["last_login_at"] else None,suspended_until=row["suspended_until"].isoformat() if row["suspended_until"] else None,warning_count=row["warning_count"] or 0,rag_input_tokens=usage.get("input_tokens"),rag_output_tokens=usage.get("output_tokens"),rag_total_tokens=usage.get("total_tokens"),rag_usage_calls=usage.get("usage_calls",0),rag_usage_missing=usage.get("usage_missing",0))


def _usage_bounds(period: str, start_date: date | None, end_date: date | None):
    zone = ZoneInfo("Asia/Ho_Chi_Minh")
    today = datetime.now(zone).date()
    if period == "custom":
        if start_date is None or end_date is None or start_date > end_date:
            raise HTTPException(400, "Cần chọn khoảng ngày hợp lệ.")
        first, last = start_date, end_date
    elif period == "all":
        return None, None
    else:
        days = {"today": 1, "7d": 7, "30d": 30}[period]
        first, last = today - timedelta(days=days-1), today
    return (datetime.combine(first, time.min, zone).astimezone(UTC),
            datetime.combine(last + timedelta(days=1), time.min, zone).astimezone(UTC))


def _usage_totals(db: Session, user_ids: list[uuid.UUID], start, end) -> dict[uuid.UUID, dict]:
    if not user_ids:
        return {}
    statement = text("""
        SELECT user_id,sum(input_tokens)::bigint input_tokens,sum(output_tokens)::bigint output_tokens,
               sum(total_tokens)::bigint total_tokens,count(*) usage_calls,
               count(*) FILTER(WHERE total_tokens IS NULL) usage_missing
        FROM rag_token_usage WHERE user_id IN :ids
          AND (:start IS NULL OR created_at >= :start) AND (:end IS NULL OR created_at < :end)
        GROUP BY user_id
    """).bindparams(bindparam("ids", expanding=True))
    return {row["user_id"]: dict(row) for row in db.execute(statement,
             {"ids": user_ids, "start": start, "end": end}).mappings()}


def users(db: Session, q: str, account_status: str | None, limit: int, offset: int,
          period: str = "30d", start_date: date | None = None,
          end_date: date | None = None, role: str | None = None) -> AdminUserPage:
    start, end = _usage_bounds(period, start_date, end_date)
    where=["u.deleted_at IS NULL"];params={"pattern":f"%{q.strip()}%","limit":limit,"offset":offset}
    if q.strip(): where.append("(u.email ILIKE :pattern OR u.username ILIKE :pattern OR p.full_name ILIKE :pattern OR p.student_code ILIKE :pattern OR p.faculty ILIKE :pattern)")
    if account_status: where.append("u.status::text=:status");params["status"]=account_status
    is_group_admin = "EXISTS (SELECT 1 FROM group_members gm WHERE gm.user_id=u.id AND gm.role='ADMIN')"
    is_admin = f"(u.system_role::text='SUPER_ADMIN' OR {is_group_admin})"
    if role == "ADMIN": where.append(is_admin)
    if role == "USER": where.append(f"NOT {is_admin}")
    clause=" AND ".join(where)
    total=_count(db,f"SELECT count(*) FROM users u LEFT JOIN profiles p ON p.user_id=u.id WHERE {clause}",params)
    rows=db.execute(text(f"SELECT u.id,u.email,u.username,u.status,u.system_role,u.created_at,u.last_login_at,u.suspended_until,u.warning_count,COALESCE(p.full_name,u.username) full_name,p.student_code,p.faculty FROM users u LEFT JOIN profiles p ON p.user_id=u.id WHERE {clause} ORDER BY u.created_at DESC LIMIT :limit OFFSET :offset"),params).mappings().all()
    groups_by_user: dict[uuid.UUID, list[dict]] = {row["id"]: [] for row in rows}
    if groups_by_user:
        group_query = text("""
            SELECT gm.user_id, g.id::text AS id, g.name, g.slug
            FROM group_members gm JOIN groups g ON g.id=gm.group_id
            WHERE gm.user_id IN :user_ids AND gm.role='ADMIN'
            ORDER BY g.name
        """).bindparams(bindparam("user_ids", expanding=True))
        for group in db.execute(group_query, {"user_ids": list(groups_by_user)}).mappings():
            groups_by_user[group["user_id"]].append({
                "id": group["id"], "name": group["name"], "slug": group["slug"],
            })
    usage = _usage_totals(db, list(groups_by_user), start, end)
    return AdminUserPage(items=[_serialize_user(row, groups_by_user[row["id"]], usage.get(row["id"])) for row in rows],total=total,limit=limit,offset=offset)


def user_rag_usage(db: Session, user_id: uuid.UUID, period: str = "30d",
                   start_date: date | None = None, end_date: date | None = None,
                   offset: int = 0) -> RAGUserUsageResponse:
    if not db.execute(text("SELECT 1 FROM users WHERE id=:uid AND deleted_at IS NULL"), {"uid": user_id}).first():
        raise HTTPException(404, "Không tìm thấy tài khoản.")
    start, end = _usage_bounds(period, start_date, end_date)
    params = {"uid": user_id, "start": start, "end": end}
    totals = _usage_totals(db, [user_id], start, end).get(user_id, {})
    daily = db.execute(text("""
        SELECT timezone('Asia/Ho_Chi_Minh',created_at)::date AS usage_day,
               sum(input_tokens)::bigint input_tokens,sum(output_tokens)::bigint output_tokens,
               sum(total_tokens)::bigint total_tokens,count(*) calls
        FROM rag_token_usage WHERE user_id=:uid
          AND (:start IS NULL OR created_at >= :start) AND (:end IS NULL OR created_at < :end)
        GROUP BY 1 ORDER BY 1
    """), params).mappings().all()
    calls = db.execute(text("""
        SELECT id,created_at,provider,model,input_tokens,output_tokens,total_tokens,status,usage_source
        FROM rag_token_usage WHERE user_id=:uid
          AND (:start IS NULL OR created_at >= :start) AND (:end IS NULL OR created_at < :end)
        ORDER BY created_at DESC,id DESC LIMIT 10 OFFSET :offset
    """), {**params, "offset": offset}).mappings().all()
    return RAGUserUsageResponse(user_id=user_id,input_tokens=totals.get("input_tokens"),
        output_tokens=totals.get("output_tokens"),total_tokens=totals.get("total_tokens"),
        usage_calls=totals.get("usage_calls",0),usage_missing=totals.get("usage_missing",0),
        daily=[{"date":row["usage_day"].isoformat(),"input_tokens":row["input_tokens"],
                "output_tokens":row["output_tokens"],"total_tokens":row["total_tokens"],
                "calls":row["calls"]} for row in daily],
        calls=[{**dict(row),"created_at":row["created_at"].isoformat()} for row in calls])


def _target(db: Session, user_id: uuid.UUID, actor: User) -> User:
    target=db.get(User,user_id)
    if not target or target.deleted_at is not None: raise HTTPException(404,"Không tìm thấy tài khoản.")
    if target.id==actor.id or target.system_role==SystemRole.SUPER_ADMIN: raise HTTPException(403,"Không thể thay đổi tài khoản Super Admin.")
    return target


def _audit(db: Session, actor: User, action: str, target: uuid.UUID, metadata: dict) -> None:
    db.execute(text("INSERT INTO audit_logs(actor_id,action,target_type,target_id,metadata) VALUES(:actor,:action,'USER',:target,CAST(:metadata AS jsonb))"),{"actor":actor.id,"action":action,"target":target,"metadata":__import__('json').dumps(metadata,ensure_ascii=False)})


def set_status(db: Session, user_id: uuid.UUID, actor: User, value: str, reason: str) -> None:
    target=_target(db,user_id,actor);old=target.status.value;target.status=AccountStatus(value)
    if target.status!=AccountStatus.ACTIVE: db.execute(text("UPDATE refresh_tokens SET revoked_at=now() WHERE user_id=:uid AND revoked_at IS NULL"),{"uid":target.id})
    _audit(db,actor,"ACCOUNT_STATUS_CHANGED",target.id,{"from":old,"to":value,"reason":reason})
    db.execute(text("INSERT INTO notifications(user_id,type,title,content,actor_id,reference_type,reference_id) VALUES(:uid,'SYSTEM','Trạng thái tài khoản thay đổi',:content,:actor,'USER',:uid)"),{"uid":target.id,"content":f"Tài khoản chuyển sang {value}. Lý do: {reason}","actor":actor.id})
    db.commit()


def set_group_admin(db: Session, user_id: uuid.UUID, group_id: uuid.UUID, actor: User, grant: bool) -> None:
    # Fail closed for legacy callers. Assignments belong to the provisioning script,
    # not the runtime Super Admin account management workflow.
    raise HTTPException(403, "Không hỗ trợ thay đổi phân quyền tài khoản trong hệ thống.")


def discipline(db:Session,user_id:uuid.UUID,actor:User,payload:DisciplineRequest)->None:
    from datetime import timedelta
    target=_target(db,user_id,actor);now=datetime.now(UTC)
    if payload.action=="WARN":
        target.warning_count+=1
    elif payload.action=="SUSPEND":
        if not payload.duration_days:raise HTTPException(400,"Phải chọn thời hạn tạm khóa.")
        target.status=AccountStatus.LOCKED;target.suspended_until=now+timedelta(days=payload.duration_days)
    elif payload.action=="BAN":
        target.status=AccountStatus.DISABLED;target.suspended_until=None
    else:
        target.status=AccountStatus.ACTIVE;target.suspended_until=None;target.failed_login_attempts=0;target.login_locked_until=None
    if payload.action in {"SUSPEND","BAN"}:db.execute(text("UPDATE refresh_tokens SET revoked_at=now() WHERE user_id=:uid AND revoked_at IS NULL"),{"uid":target.id})
    _audit(db,actor,f"USER_{payload.action}",target.id,{"reason":payload.reason,"duration_days":payload.duration_days})
    db.execute(text("INSERT INTO notifications(user_id,type,title,content,actor_id,reference_type,reference_id) VALUES(:uid,'SYSTEM','Xử lý tài khoản',:content,:actor,'USER',:uid)"),{"uid":target.id,"content":f"{payload.action}: {payload.reason}","actor":actor.id})
    db.commit()


def audit_logs(db:Session,limit:int=100)->list[AuditLogResponse]:
    rows=db.execute(text("SELECT a.id,a.action,a.target_type,a.target_id,a.metadata,a.created_at,COALESCE(p.full_name,u.username,'Hệ thống') actor_name FROM audit_logs a LEFT JOIN users u ON u.id=a.actor_id LEFT JOIN profiles p ON p.user_id=u.id ORDER BY a.created_at DESC LIMIT :limit"),{"limit":limit}).mappings()
    return [AuditLogResponse(id=str(r["id"]),actor_name=r["actor_name"],action=r["action"],target_type=r["target_type"],target_id=str(r["target_id"]) if r["target_id"] else None,metadata=r["metadata"] or {},created_at=r["created_at"].isoformat()) for r in rows]
