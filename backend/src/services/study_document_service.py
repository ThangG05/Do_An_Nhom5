"""Private, group-scoped document library for the canonical Study group only."""
import hashlib
import logging
import tempfile
import uuid
import zipfile
from collections import defaultdict
from pathlib import Path
from urllib.parse import quote

from botocore.exceptions import BotoCoreError, ClientError
from fastapi import HTTPException, UploadFile
from PIL import Image, UnidentifiedImageError
from docx import Document as WordDocument
from openpyxl import load_workbook
from sqlalchemy import bindparam, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from src.config import settings
from src.db.models.user import SystemRole, User
from src.models.study_document import (StudyDocumentFileResponse, StudyDocumentResponse,
                                       SubjectResponse)
from src.models.study_document_report import (StudyDocumentReportResponse,
                                               StudyDocumentReportDecision)
from src.services.storage_service import get_r2_client

logger = logging.getLogger(__name__)
MAX_SHARE_BYTES = 50 * 1024 * 1024
MAX_FILES = 10
MIME_BY_EXTENSION = {
    "pdf": "application/pdf",
    "doc": "application/msword",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "xls": "application/vnd.ms-excel",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "ppt": "application/vnd.ms-powerpoint",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "jpg": "image/jpeg", "jpeg": "image/jpeg", "png": "image/png", "webp": "image/webp",
}
OOXML_PARTS = {"docx": "word/", "xlsx": "xl/", "pptx": "ppt/"}
IMAGE_FORMATS = {"jpg": "JPEG", "jpeg": "JPEG", "png": "PNG", "webp": "WEBP"}


def _private_bucket() -> str:
    bucket = settings.R2_STUDY_BUCKET_NAME.strip()
    if not bucket or bucket == settings.R2_BUCKET_NAME.strip():
        raise HTTPException(503, "Kho tài liệu cần bucket R2 riêng tư, khác bucket media công khai.")
    return bucket


def require_study_group(db: Session, group_id: uuid.UUID) -> None:
    slug = db.execute(text("SELECT slug FROM groups WHERE id=:id AND status='ACTIVE'"),
                      {"id": group_id}).scalar_one_or_none()
    if slug != "hoc-tap":
        raise HTTPException(404, "Kho tài liệu chỉ có trong nhóm Học tập.")


def require_member(db: Session, group_id: uuid.UUID, user: User) -> None:
    require_study_group(db, group_id)
    if user.system_role == SystemRole.SUPER_ADMIN:
        return
    role = db.execute(text("SELECT role FROM group_members WHERE group_id=:gid AND user_id=:uid"),
                      {"gid": group_id, "uid": user.id}).scalar_one_or_none()
    if str(role) not in {"MEMBER", "ADMIN"}:
        raise HTTPException(403, "Bạn cần là thành viên nhóm Học tập.")


def require_admin(db: Session, group_id: uuid.UUID, user: User) -> None:
    require_study_group(db, group_id)
    if user.system_role == SystemRole.SUPER_ADMIN:
        return
    role = db.execute(text("SELECT role FROM group_members WHERE group_id=:gid AND user_id=:uid"),
                      {"gid": group_id, "uid": user.id}).scalar_one_or_none()
    if str(role) != "ADMIN":
        raise HTTPException(403, "Chỉ admin của nhóm Học tập được duyệt tài liệu.")


def require_subject_admin(db: Session, group_id: uuid.UUID, user: User) -> None:
    """Only the assigned Học tập group admin may change the shared subject catalog."""
    require_study_group(db, group_id)
    if user.system_role == SystemRole.SUPER_ADMIN:
        raise HTTPException(403, "Danh mục môn học thuộc quyền quản trị nhóm Học tập.")
    role = db.execute(text("SELECT role FROM group_members WHERE group_id=:gid AND user_id=:uid"),
                      {"gid": group_id, "uid": user.id}).scalar_one_or_none()
    if str(role) != "ADMIN":
        raise HTTPException(403, "Chỉ admin của nhóm Học tập được quản lý môn học.")


def list_subjects(db: Session, query: str = "", *, include_inactive: bool = False,
                  category: str | None = None, major_id: uuid.UUID | None = None,
                  limit: int = 50, offset: int = 0) -> list[SubjectResponse]:
    query = query.strip()
    tokens = [token for token in query.split() if token][:8]
    token_match = " AND ".join(
        f"(s.name ILIKE :token_{index} OR s.code ILIKE :token_{index})"
        for index, _ in enumerate(tokens)
    ) or "FALSE"
    params = {"all": include_inactive, "q": query, "pattern": f"%{query}%",
              "prefix": f"{query}%", "category": category, "major": major_id,
              "limit": limit, "offset": offset}
    params.update({f"token_{index}": f"%{token}%" for index, token in enumerate(tokens)})
    rows = db.execute(text(f"""
        SELECT s.*, ARRAY(SELECT ms.major_id FROM major_subjects ms WHERE ms.subject_id=s.id) major_ids
        FROM subjects s
        WHERE (:all OR s.is_active) AND (:q='' OR s.code ILIKE :pattern OR s.name ILIKE :pattern
              OR ({token_match}))
          AND (:category IS NULL OR s.category=:category)
          AND (CAST(:major AS uuid) IS NULL OR EXISTS (
              SELECT 1 FROM major_subjects ms WHERE ms.subject_id=s.id AND ms.major_id=:major))
        ORDER BY CASE
          WHEN :q='' THEN 9
          WHEN lower(s.code)=lower(:q) THEN 0
          WHEN s.code ILIKE :prefix THEN 1
          WHEN s.name ILIKE :prefix THEN 2
          WHEN s.code ILIKE :pattern THEN 3
          WHEN s.name ILIKE :pattern THEN 4
          ELSE 5 END, s.name,s.code,s.id
        LIMIT :limit OFFSET :offset
    """), params).mappings()
    return [SubjectResponse(**dict(row)) for row in rows]


def list_majors(db: Session) -> list[dict]:
    return [dict(row) for row in db.execute(text("""
        SELECT m.id,m.code,m.name,count(s.id) subject_count FROM majors m
        LEFT JOIN major_subjects ms ON ms.major_id=m.id
        LEFT JOIN subjects s ON s.id=ms.subject_id AND s.is_active
        GROUP BY m.id ORDER BY m.name
    """)).mappings()]


def upsert_subject(db: Session, code: str | None, name: str, source_note: str, is_active: bool,
                   *, subject_id: uuid.UUID | None = None, category: str = "SPECIALIZED",
                   needs_review: bool = False, major_ids: list[uuid.UUID] | None = None) -> SubjectResponse:
    code = code.strip().upper() if code and code.strip() else None
    name, source_note = name.strip(), source_note.strip()
    if len(name) < 2 or len(source_note) < 5:
        raise HTTPException(400, "Cần tên môn và nguồn đối chiếu đầy đủ.")
    if category == "GENERAL" and major_ids:
        raise HTTPException(400, "Môn đại cương dùng chung, không gán ngành bắt buộc.")
    if major_ids:
        found = db.execute(text("SELECT id FROM majors WHERE id IN :ids").bindparams(
            bindparam("ids", expanding=True)), {"ids": major_ids}).scalars().all()
        if set(found) != set(major_ids):
            raise HTTPException(400, "Ngành được chọn không có trong danh mục.")
    params = {"code": code, "name": name, "source": source_note, "active": is_active,
              "category": category, "review": needs_review or code is None, "id": subject_id or uuid.uuid4()}
    try:
        if subject_id:
            row = db.execute(text("""UPDATE subjects SET code=:code,name=:name,source_note=:source,
                is_active=:active,category=:category,needs_review=:review,updated_at=now()
                WHERE id=:id RETURNING *"""), params).mappings().first()
            if not row:
                raise HTTPException(404, "Không tìm thấy môn học.")
        else:
            row = db.execute(text("""INSERT INTO subjects(id,code,name,source_note,is_active,category,needs_review)
                VALUES(:id,:code,:name,:source,:active,:category,:review) RETURNING *"""), params).mappings().one()
        sid = row["id"]
        if code:
            conflict = db.execute(text("SELECT count(DISTINCT name)>1 FROM subjects WHERE code=:code"),
                                  {"code": code}).scalar()
            if conflict:
                db.execute(text("UPDATE subjects SET needs_review=true WHERE code=:code"), {"code": code})
                row = {**row, "needs_review": True}
        if major_ids is not None or category == "GENERAL":
            db.execute(text("DELETE FROM major_subjects WHERE subject_id=:id"), {"id": sid})
            if major_ids:
                db.execute(text("INSERT INTO major_subjects(major_id,subject_id) VALUES(:major,:subject)"),
                           [{"major": mid, "subject": sid} for mid in set(major_ids)])
        ids = db.execute(text("SELECT major_id FROM major_subjects WHERE subject_id=:id"), {"id": sid}).scalars().all()
        result = SubjectResponse(**dict(row), major_ids=ids)
        db.commit()
        return result
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "Môn học cùng mã, tên và phân loại đã tồn tại. Hãy chọn Chỉnh sửa bản ghi đó.")


def _validate_file(file: UploadFile) -> dict:
    name = Path((file.filename or "").replace("\\", "/")).name
    name = "".join(character for character in name if character.isprintable() and character not in '\"\\')
    extension = Path(name).suffix.lower().lstrip(".")
    mime = MIME_BY_EXTENSION.get(extension)
    if not name or not mime:
        raise HTTPException(400, "Chỉ nhận PDF, Word, Excel, PowerPoint và ảnh JPG/PNG/WEBP.")
    declared = (file.content_type or "").split(";", 1)[0].lower()
    if declared not in {mime, "application/octet-stream"}:
        raise HTTPException(400, f"MIME của {name} không khớp phần mở rộng.")
    stream = file.file
    stream.seek(0, 2)
    size = stream.tell()
    stream.seek(0)
    if not 0 < size <= MAX_SHARE_BYTES:
        raise HTTPException(400, "File trống hoặc vượt giới hạn 50 MB.")
    prefix = stream.read(16)
    stream.seek(0)
    if extension == "pdf":
        valid = prefix.startswith(b"%PDF-")
    elif extension in {"doc", "xls", "ppt"}:
        valid = prefix.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    elif extension in OOXML_PARTS:
        valid = False
        if prefix.startswith(b"PK"):
            try:
                with zipfile.ZipFile(stream) as archive:
                    names = archive.namelist()
                    valid = "[Content_Types].xml" in names and any(
                        part.startswith(OOXML_PARTS[extension]) for part in names)
            except (OSError, zipfile.BadZipFile):
                pass
    else:
        try:
            with Image.open(stream) as image:
                valid = image.format == IMAGE_FORMATS[extension]
                image.verify()
        except (OSError, UnidentifiedImageError):
            valid = False
    stream.seek(0)
    if not valid:
        raise HTTPException(400, f"Nội dung thực tế của {name} không đúng định dạng.")
    digest = hashlib.sha256()
    while chunk := stream.read(1024 * 1024):
        digest.update(chunk)
    stream.seek(0)
    return {"name": name[:255], "mime": mime, "extension": extension,
            "size": size, "sha256": digest.hexdigest(), "stream": stream}


def _serialize(db: Session, rows: list) -> list[StudyDocumentResponse]:
    if not rows:
        return []
    statement = text("""
        SELECT id,document_id,original_name,mime_type,extension,size_bytes
        FROM study_document_files WHERE document_id IN :ids AND deleted_at IS NULL
        ORDER BY created_at,id
    """).bindparams(bindparam("ids", expanding=True))
    files: dict[uuid.UUID, list[StudyDocumentFileResponse]] = defaultdict(list)
    for row in db.execute(statement, {"ids": [item["id"] for item in rows]}).mappings():
        files[row["document_id"]].append(StudyDocumentFileResponse(**{
            key: row[key] for key in ("id", "original_name", "mime_type", "extension", "size_bytes")
        }))
    return [StudyDocumentResponse(**{key: value for key, value in row.items()
                                     if key not in {"created_at", "reviewed_at"}}, files=files[row["id"]],
            created_at=row["created_at"].isoformat(),
            reviewed_at=row["reviewed_at"].isoformat() if row["reviewed_at"] else None)
            for row in rows]


_DETAIL_SQL = """
    SELECT d.id,d.group_id,d.subject_id,s.code subject_code,s.name subject_name,
           d.uploaded_by,COALESCE(p.full_name,u.username) uploader_name,d.title,d.description,
           d.status,d.total_bytes,d.reviewed_at,d.rejection_reason,d.created_at
    FROM study_documents d JOIN subjects s ON s.id=d.subject_id
    JOIN users u ON u.id=d.uploaded_by LEFT JOIN profiles p ON p.user_id=u.id
"""


def list_documents(db: Session, group_id: uuid.UUID, user: User, *, mine: bool = False,
                   pending: bool = False, query: str = "", subject_id: uuid.UUID | None = None,
                   limit: int = 30, offset: int = 0, category: str | None = None,
                   major_id: uuid.UUID | None = None) -> list[StudyDocumentResponse]:
    (require_admin if pending else require_member)(db, group_id, user)
    condition = "d.status='PENDING'" if pending else "d.uploaded_by=:uid" if mine else "d.status='APPROVED'"
    rows = db.execute(text(_DETAIL_SQL + f"""
        WHERE d.group_id=:gid AND d.deleted_at IS NULL AND {condition}
          AND (:subject_id IS NULL OR d.subject_id=:subject_id)
          AND (:category IS NULL OR s.category=:category)
          AND (CAST(:major AS uuid) IS NULL OR EXISTS
               (SELECT 1 FROM major_subjects ms WHERE ms.subject_id=s.id AND ms.major_id=:major))
          AND (:q='' OR d.title ILIKE :pattern OR s.name ILIKE :pattern OR s.code ILIKE :pattern)
        ORDER BY d.created_at DESC LIMIT :limit OFFSET :offset
    """), {"gid": group_id, "uid": user.id, "subject_id": subject_id, "q": query.strip(),
             "pattern": f"%{query.strip()}%", "limit": limit, "offset": offset,
             "category": category, "major": major_id}).mappings().all()
    return _serialize(db, rows)


def count_pending_documents(db: Session, group_id: uuid.UUID, admin: User) -> int:
    require_admin(db, group_id, admin)
    return int(db.execute(text("""
        SELECT count(*) FROM study_documents
        WHERE group_id=:gid AND status='PENDING' AND deleted_at IS NULL
    """), {"gid": group_id}).scalar_one())


def count_open_document_reports(db: Session, group_id: uuid.UUID, admin: User) -> int:
    require_admin(db, group_id, admin)
    return int(db.execute(text("""
        SELECT count(*) FROM study_document_reports
        WHERE group_id=:gid AND status='OPEN'
    """), {"gid": group_id}).scalar_one())


def create_document_report(db: Session, group_id: uuid.UUID, document_id: uuid.UUID,
                           user: User, reason: str) -> dict:
    require_member(db, group_id, user)
    reason = reason.strip()
    if len(reason) < 5:
        raise HTTPException(400, "Vui lòng ghi rõ nguyên nhân tố cáo (ít nhất 5 ký tự).")
    document = db.execute(text("""
        SELECT d.uploaded_by, COALESCE(p.full_name,u.username) reporter_name
        FROM study_documents d JOIN users u ON u.id=:uid
        LEFT JOIN profiles p ON p.user_id=u.id
        WHERE d.id=:did AND d.group_id=:gid AND d.status='APPROVED' AND d.deleted_at IS NULL
    """), {"did": document_id, "gid": group_id, "uid": user.id}).mappings().first()
    if not document:
        raise HTTPException(404, "Không tìm thấy tài liệu đã được duyệt.")
    if document["uploaded_by"] == user.id:
        raise HTTPException(400, "Bạn không thể tố cáo tài liệu do chính mình chia sẻ.")
    report_id = uuid.uuid4()
    try:
        db.execute(text("""
            INSERT INTO study_document_reports(id,group_id,document_id,reporter_id,reporter_name,reason)
            VALUES(:id,:gid,:did,:uid,:name,:reason)
        """), {"id": report_id, "gid": group_id, "did": document_id, "uid": user.id,
                 "name": document["reporter_name"], "reason": reason})
        db.commit()
    except IntegrityError as exc:
        db.rollback()
        raise HTTPException(409, "Bạn đã gửi tố cáo tài liệu này và admin đang xem xét.") from exc
    return {"id": report_id, "status": "OPEN"}


def list_open_document_reports(db: Session, group_id: uuid.UUID, admin: User,
                               limit: int = 50, offset: int = 0) -> list[StudyDocumentReportResponse]:
    require_admin(db, group_id, admin)
    rows = db.execute(text("""
        SELECT r.id,r.group_id,r.document_id,d.title,s.code subject_code,s.name subject_name,
               COALESCE(owner_profile.full_name,owner.username,'Tài khoản đã xóa') uploader_name,
               r.reporter_name,r.reason,r.status,r.resolution_note,
               COALESCE(handler_profile.full_name,handler.username) handled_by_name,r.created_at,
               (d.deleted_at IS NOT NULL) document_deleted
        FROM study_document_reports r
        JOIN study_documents d ON d.id=r.document_id
        JOIN subjects s ON s.id=d.subject_id
        LEFT JOIN users owner ON owner.id=d.uploaded_by
        LEFT JOIN profiles owner_profile ON owner_profile.user_id=owner.id
        LEFT JOIN users handler ON handler.id=r.handled_by
        LEFT JOIN profiles handler_profile ON handler_profile.user_id=handler.id
        WHERE r.group_id=:gid AND r.status='OPEN'
        ORDER BY r.created_at ASC,r.id LIMIT :limit OFFSET :offset
    """), {"gid": group_id, "limit": limit, "offset": offset}).mappings().all()
    return [StudyDocumentReportResponse(**dict(row)) for row in rows]


def decide_document_report(db: Session, group_id: uuid.UUID, report_id: uuid.UUID,
                           admin: User, payload: StudyDocumentReportDecision) -> dict:
    require_admin(db, group_id, admin)
    note = payload.note.strip()
    if len(note) < 3:
        raise HTTPException(400, "Vui lòng ghi lý do xử lý (ít nhất 3 ký tự).")
    report = db.execute(text("""
        SELECT r.document_id,d.deleted_at,d.uploaded_by,d.title,r.reporter_id FROM study_document_reports r
        JOIN study_documents d ON d.id=r.document_id
        WHERE r.id=:rid AND r.group_id=:gid AND r.status='OPEN'
        FOR UPDATE OF r,d
    """), {"rid": report_id, "gid": group_id}).mappings().first()
    if not report:
        raise HTTPException(404, "Tố cáo không còn trong hàng chờ xử lý.")
    document_id = report["document_id"]
    if payload.decision == "REMOVE_DOCUMENT" and report["deleted_at"] is None:
        files = db.execute(text("""
            SELECT id,r2_object_key FROM study_document_files
            WHERE document_id=:did AND deleted_at IS NULL
        """), {"did": document_id}).mappings().all()
        client, bucket = get_r2_client(), _private_bucket()
        try:
            for file in files:
                client.delete_object(Bucket=bucket, Key=file["r2_object_key"])
        except (BotoCoreError, ClientError) as exc:
            db.rollback()
            logger.warning("study_document_report_r2_remove_failed")
            raise HTTPException(502, "Chưa thể gỡ file khỏi kho lưu trữ. Tố cáo vẫn được giữ để admin thử lại.") from exc
        db.execute(text("UPDATE study_document_files SET deleted_at=now() WHERE document_id=:did AND deleted_at IS NULL"),
                   {"did": document_id})
        db.execute(text("UPDATE study_documents SET deleted_at=now() WHERE id=:did AND deleted_at IS NULL"),
                   {"did": document_id})
    status = "RESOLVED" if payload.decision == "REMOVE_DOCUMENT" else "DISMISSED"
    db.execute(text("""
        UPDATE study_document_reports SET status=:status,resolution_note=:note,handled_by=:admin,handled_at=now()
        WHERE id=:rid
    """), {"status": status, "note": note, "admin": admin.id, "rid": report_id})
    if payload.decision == "REMOVE_DOCUMENT":
        db.execute(text("""
            INSERT INTO notifications(user_id,type,title,content,actor_id,reference_type,reference_id)
            VALUES(:uid,'SYSTEM','Tài liệu bị gỡ khỏi kho Học tập',:content,:actor,'GROUP',:gid)
        """), {"uid": report["uploaded_by"], "content": f"Admin đã gỡ tài liệu '{report['title']}' khỏi kho Học tập. Lý do: {note}",
                 "actor": admin.id, "gid": group_id})
    if report["reporter_id"]:
        outcome = "đã được admin xử lý" if payload.decision == "REMOVE_DOCUMENT" else "đã được admin xem xét và bác bỏ"
        db.execute(text("""
            INSERT INTO notifications(user_id,type,title,content,actor_id,reference_type,reference_id)
            VALUES(:uid,'SYSTEM','Cập nhật tố cáo tài liệu',:content,:actor,'GROUP',:gid)
        """), {"uid": report["reporter_id"], "content": f"Tố cáo của bạn {outcome}. Lý do xử lý: {note}",
                 "actor": admin.id, "gid": group_id})
    db.commit()
    return {"id": report_id, "status": status, "document_id": document_id}


def get_document(db: Session, group_id: uuid.UUID, document_id: uuid.UUID, user: User,
                 *, for_admin: bool = False) -> StudyDocumentResponse:
    (require_admin if for_admin else require_member)(db, group_id, user)
    row = db.execute(text(_DETAIL_SQL + " WHERE d.id=:id AND d.group_id=:gid AND d.deleted_at IS NULL"),
                     {"id": document_id, "gid": group_id}).mappings().first()
    if not row or (not for_admin and row["status"] != "APPROVED" and row["uploaded_by"] != user.id):
        raise HTTPException(404, "Không tìm thấy tài liệu.")
    return _serialize(db, [row])[0]


def create_document(db: Session, group_id: uuid.UUID, user: User, subject_id: uuid.UUID,
                    title: str, description: str | None, files: list[UploadFile]) -> StudyDocumentResponse:
    require_member(db, group_id, user)
    if not title.strip() or not files or len(files) > MAX_FILES:
        raise HTTPException(400, "Cần tên tài liệu và từ 1 đến 10 file.")
    subject = db.execute(text("SELECT 1 FROM subjects WHERE id=:id AND is_active"),
                         {"id": subject_id}).first()
    if not subject:
        raise HTTPException(400, "Môn học không có trong danh mục đang sử dụng.")
    validated = [_validate_file(file) for file in files]
    if len({item["sha256"] for item in validated}) != len(validated):
        raise HTTPException(400, "Một lượt chia sẻ không thể đính kèm hai file có nội dung trùng nhau.")
    total = sum(item["size"] for item in validated)
    if total > MAX_SHARE_BYTES:
        raise HTTPException(400, "Tổng dung lượng một lần chia sẻ không được vượt 50 MB.")
    document_id = uuid.uuid4()
    uploaded: list[str] = []
    bucket = _private_bucket()
    client = get_r2_client()
    try:
        for item in validated:
            item["id"] = uuid.uuid4()
            item["key"] = f"study-groups/{group_id}/documents/{document_id}/{item['id']}.{item['extension']}"
            # Register the key before starting transfer: even a partial/failed upload may leave an object.
            uploaded.append(item["key"])
            client.upload_fileobj(item["stream"], bucket, item["key"],
                                  ExtraArgs={"ContentType": item["mime"]})
        db.execute(text("""
            INSERT INTO study_documents(id,group_id,subject_id,uploaded_by,title,description,status,total_bytes,file_count)
            VALUES(:id,:gid,:sid,:uid,:title,:description,'PENDING',:total,:count)
        """), {"id": document_id, "gid": group_id, "sid": subject_id, "uid": user.id,
               "title": title.strip(), "description": description.strip() if description else None,
               "total": total, "count": len(validated)})
        for item in validated:
            db.execute(text("""
                INSERT INTO study_document_files(id,document_id,r2_object_key,original_name,mime_type,
                                                 extension,size_bytes,checksum_sha256)
                VALUES(:id,:did,:key,:name,:mime,:extension,:size,:sha256)
            """), {"did": document_id, **{key: item[key] for key in
                    ("id", "key", "name", "mime", "extension", "size", "sha256")}})
        db.commit()
    except Exception as exc:
        db.rollback()
        for key in uploaded:
            try:
                client.delete_object(Bucket=bucket, Key=key)
            except (BotoCoreError, ClientError):
                logger.warning("study_upload_orphan_cleanup_failed")
        if isinstance(exc, (BotoCoreError, ClientError)):
            error_code = (
                exc.response.get("Error", {}).get("Code", "unknown")
                if isinstance(exc, ClientError) else type(exc).__name__
            )
            logger.warning("study_document_r2_upload_failed", extra={"r2_error_code": error_code})
            raise HTTPException(
                status_code=502,
                detail="Không thể tải tài liệu lên kho R2. Hãy kiểm tra quyền đọc/ghi của API token với bucket tài liệu.",
            ) from exc
        raise
    return get_document(db, group_id, document_id, user)


def file_download(db: Session, group_id: uuid.UUID, document_id: uuid.UUID,
                  file_id: uuid.UUID, user: User, *, preview: bool = False) -> str:
    require_study_group(db, group_id)
    row = db.execute(text("""
        SELECT f.r2_object_key,f.original_name,f.mime_type,d.status,d.uploaded_by
        FROM study_document_files f JOIN study_documents d ON d.id=f.document_id
        WHERE f.id=:fid AND f.document_id=:did AND d.group_id=:gid
          AND d.deleted_at IS NULL AND f.deleted_at IS NULL
    """), {"fid": file_id, "did": document_id, "gid": group_id}).mappings().first()
    membership = db.execute(text("SELECT role FROM group_members WHERE group_id=:gid AND user_id=:uid"),
                            {"gid": group_id, "uid": user.id}).scalar_one_or_none()
    is_admin = str(membership) == "ADMIN" or user.system_role == SystemRole.SUPER_ADMIN
    if not row or (str(membership) not in {"MEMBER", "ADMIN"} and not is_admin) or (
            row["status"] != "APPROVED" and row["uploaded_by"] != user.id and not is_admin):
        raise HTTPException(404, "Không tìm thấy file.")
    if preview and row["mime_type"] not in {"application/pdf", "image/jpeg", "image/png", "image/webp"}:
        raise HTTPException(400, "Chỉ hỗ trợ xem trước PDF và ảnh; file Office cần tải về.")
    filename = quote(row["original_name"], safe="")
    disposition = "inline" if preview else "attachment"
    return get_r2_client().generate_presigned_url("get_object", Params={
        "Bucket": _private_bucket(), "Key": row["r2_object_key"],
        "ResponseContentDisposition": f"{disposition}; filename*=UTF-8''{filename}",
    }, ExpiresIn=min(settings.R2_PRESIGNED_URL_EXPIRE, 300))


def _get_document_files_for_user(db: Session, group_id: uuid.UUID, document_id: uuid.UUID,
                                 user: User) -> list:
    require_member(db, group_id, user)
    row = db.execute(text("""
        SELECT id,status,uploaded_by FROM study_documents
        WHERE id=:did AND group_id=:gid AND deleted_at IS NULL
    """), {"did": document_id, "gid": group_id}).mappings().first()
    if not row or (row["status"] != "APPROVED" and row["uploaded_by"] != user.id):
        raise HTTPException(404, "Không tìm thấy tài liệu.")
    files = db.execute(text("""
        SELECT id,r2_object_key,original_name,mime_type,extension,size_bytes
        FROM study_document_files WHERE document_id=:did AND deleted_at IS NULL
        ORDER BY created_at,id
    """), {"did": document_id}).mappings().all()
    if not files:
        raise HTTPException(404, "Bộ tài liệu hiện không có file để tải.")
    return files


def download_document_bundle(db: Session, group_id: uuid.UUID, document_id: uuid.UUID,
                             user: User) -> tuple[tempfile.SpooledTemporaryFile, str]:
    files = _get_document_files_for_user(db, group_id, document_id, user)
    archive_file = tempfile.SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    client = get_r2_client()
    try:
        with zipfile.ZipFile(archive_file, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            for index, file in enumerate(files, start=1):
                filename = Path(file["original_name"].replace("\\", "/")).name or f"file-{index}"
                archive_name = f"{index:02d} - {filename}"
                with archive.open(archive_name, "w") as target:
                    client.download_fileobj(_private_bucket(), file["r2_object_key"], target)
        archive_file.seek(0)
        return archive_file, f"bo-tai-lieu-{document_id}.zip"
    except (BotoCoreError, ClientError) as exc:
        archive_file.close()
        logger.warning("study_document_bundle_download_failed")
        raise HTTPException(502, "Không thể tải trọn bộ tài liệu từ kho R2.") from exc
    except Exception:
        archive_file.close()
        raise


def preview_office_file(db: Session, group_id: uuid.UUID, document_id: uuid.UUID,
                        file_id: uuid.UUID, user: User) -> dict:
    files = _get_document_files_for_user(db, group_id, document_id, user)
    file = next((item for item in files if item["id"] == file_id), None)
    if not file:
        raise HTTPException(404, "Không tìm thấy file.")
    extension = file["extension"].lower()
    if extension not in {"docx", "xlsx"}:
        raise HTTPException(415, "Xem trước hiện hỗ trợ Word .docx và Excel .xlsx; file .doc/.xls cần tải về.")
    content = tempfile.SpooledTemporaryFile(max_size=8 * 1024 * 1024, mode="w+b")
    try:
        get_r2_client().download_fileobj(_private_bucket(), file["r2_object_key"], content)
    except (BotoCoreError, ClientError) as exc:
        content.close()
        logger.warning("study_document_office_preview_download_failed")
        raise HTTPException(502, "Không thể đọc file xem trước từ kho R2.") from exc
    content.seek(0)
    try:
        if extension == "docx":
            document = WordDocument(content)
            paragraphs = [paragraph.text[:10_000] for paragraph in document.paragraphs if paragraph.text.strip()]
            tables = [[[cell.text[:5_000] for cell in row.cells[:30]] for row in table.rows[:200]]
                      for table in document.tables[:50]]
            return {"kind": "word", "name": file["original_name"],
                    "paragraphs": paragraphs[:500], "tables": tables[:50]}

        workbook = load_workbook(content, read_only=True, data_only=True)
        try:
            sheets = []
            for sheet in workbook.worksheets[:20]:
                rows = []
                for row in sheet.iter_rows(max_row=201, max_col=30, values_only=True):
                    values = [value.isoformat() if hasattr(value, "isoformat") else
                              str(value) if value is not None else None for value in row]
                    if any(value is not None for value in values):
                        rows.append(values)
                sheets.append({"name": sheet.title, "rows": rows[:200],
                               "truncated": sheet.max_row > 200 or sheet.max_column > 30})
        finally:
            workbook.close()
        return {"kind": "excel", "name": file["original_name"], "sheets": sheets}
    except Exception as exc:
        raise HTTPException(422, "File Office không đọc được hoặc không đúng định dạng.") from exc
    finally:
        content.close()


def delete_own_document(db: Session, group_id: uuid.UUID, document_id: uuid.UUID,
                        user: User) -> None:
    require_member(db, group_id, user)
    document = db.execute(text("""
        SELECT id FROM study_documents
        WHERE id=:did AND group_id=:gid AND uploaded_by=:uid AND deleted_at IS NULL
    """), {"did": document_id, "gid": group_id, "uid": user.id}).first()
    if not document:
        raise HTTPException(404, "Không tìm thấy tài liệu của bạn.")
    files = db.execute(text("""
        SELECT id,r2_object_key FROM study_document_files
        WHERE document_id=:did AND deleted_at IS NULL
    """), {"did": document_id}).mappings().all()
    client = get_r2_client()
    bucket = _private_bucket()
    try:
        for file in files:
            client.delete_object(Bucket=bucket, Key=file["r2_object_key"])
    except (BotoCoreError, ClientError) as exc:
        logger.warning("study_document_owner_delete_r2_failed")
        raise HTTPException(502, "Chưa thể xóa tài liệu khỏi kho R2. Vui lòng thử lại sau.") from exc
    db.execute(text("UPDATE study_document_files SET deleted_at=now() WHERE document_id=:did AND deleted_at IS NULL"),
               {"did": document_id})
    db.execute(text("UPDATE study_documents SET deleted_at=now() WHERE id=:did AND deleted_at IS NULL"),
               {"did": document_id})
    db.commit()


def moderate_document(db: Session, group_id: uuid.UUID, document_id: uuid.UUID,
                      admin: User, approve: bool, reason: str | None) -> StudyDocumentResponse:
    require_admin(db, group_id, admin)
    if not approve and (not reason or len(reason.strip()) < 3):
        raise HTTPException(400, "Cần nhập lý do từ chối.")
    row = db.execute(text("""
        UPDATE study_documents SET status=:status,reviewed_by=:admin,reviewed_at=now(),rejection_reason=:reason
        WHERE id=:id AND group_id=:gid AND status='PENDING' AND deleted_at IS NULL
        RETURNING uploaded_by
    """), {"status": "APPROVED" if approve else "REJECTED", "admin": admin.id,
             "reason": None if approve else reason.strip(), "id": document_id, "gid": group_id}).first()
    if not row:
        db.rollback()
        raise HTTPException(409, "Tài liệu không còn chờ duyệt.")
    db.execute(text("""
        INSERT INTO notifications(user_id,type,title,content,actor_id,reference_type,reference_id)
        VALUES(:uid,'SYSTEM','Kết quả duyệt tài liệu',:content,:actor,'GROUP',:gid)
    """), {"uid": row.uploaded_by, "actor": admin.id, "gid": group_id,
             "content": "Tài liệu đã được duyệt." if approve else f"Tài liệu bị từ chối: {reason.strip()}"})
    db.commit()
    # Rejected binary objects are removed immediately; metadata and the reason remain visible to the owner.
    if not approve:
        files = db.execute(text("SELECT id,r2_object_key FROM study_document_files WHERE document_id=:id AND deleted_at IS NULL"),
                           {"id": document_id}).mappings().all()
        for file in files:
            try:
                get_r2_client().delete_object(Bucket=_private_bucket(), Key=file["r2_object_key"])
                db.execute(text("UPDATE study_document_files SET deleted_at=now() WHERE id=:id"), {"id": file["id"]})
                db.commit()
            except (BotoCoreError, ClientError):
                db.rollback()
                logger.warning("study_rejected_file_cleanup_failed")
    return get_document(db, group_id, document_id, admin, for_admin=True)


def cleanup_rejected_files(db: Session, group_id: uuid.UUID, admin: User,
                           limit: int = 50) -> dict[str, int]:
    """Retry R2 cleanup if deletion failed after a committed rejection."""
    require_admin(db, group_id, admin)
    files = db.execute(text("""
        SELECT f.id,f.r2_object_key FROM study_document_files f
        JOIN study_documents d ON d.id=f.document_id
        WHERE d.group_id=:gid AND d.status='REJECTED' AND f.deleted_at IS NULL
        ORDER BY d.reviewed_at LIMIT :limit
    """), {"gid": group_id, "limit": limit}).mappings().all()
    removed = failed = 0
    bucket = _private_bucket()
    client = get_r2_client()
    for file in files:
        try:
            client.delete_object(Bucket=bucket, Key=file["r2_object_key"])
            db.execute(text("UPDATE study_document_files SET deleted_at=now() WHERE id=:id"),
                       {"id": file["id"]})
            db.commit()
            removed += 1
        except (BotoCoreError, ClientError):
            db.rollback()
            failed += 1
            logger.warning("study_rejected_file_cleanup_failed")
    return {"removed": removed, "failed": failed}
